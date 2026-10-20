# -*- coding: utf-8 -*-
import json
from pathlib import Path
from html import escape
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill

src = Path(r"C:\Users\LENOVO\Desktop\chinese-braille-toolkit\data\synthetic_tables\tables.jsonl")
out_dir = src.parent
tables = []
with src.open(encoding="utf-8") as f:
    for line in f:
        line = line.strip()
        if line:
            tables.append(json.loads(line))

parts = [
    "<!DOCTYPE html>",
    "<html lang='zh-CN'><head><meta charset='utf-8'>",
    "<title>synthetic_tables 表格预览</title>",
    "<style>",
    "body{font-family:'Segoe UI','Microsoft YaHei',sans-serif;margin:24px;background:#f6f7fb;color:#1a1a1a;}",
    "h1{font-size:1.4rem;} h2{font-size:1.15rem;margin-top:2rem;}",
    ".meta{color:#555;font-size:0.9rem;margin-bottom:8px;}",
    "table{border-collapse:collapse;background:#fff;margin:8px 0 16px;box-shadow:0 1px 3px rgba(0,0,0,.08);}",
    "th,td{border:1px solid #ccc;padding:8px 12px;text-align:left;vertical-align:top;}",
    "th{background:#e8eef8;font-weight:600;}",
    "tr:nth-child(even) td{background:#fafbfe;}",
    ".braille{font-size:1.05rem;letter-spacing:0.05em;}",
    ".pair{display:grid;grid-template-columns:1fr 1fr;gap:16px;}",
    "@media (max-width:900px){.pair{grid-template-columns:1fr;}}",
    "nav a{margin-right:12px;}",
    "</style></head><body>",
    "<h1>测评集表格预览（synthetic_tables）</h1>",
    f"<p class='meta'>共 {len(tables)} 张表 · 来源 tables.jsonl · 每张表同时展示汉字与盲文</p>",
    "<nav>" + " ".join(f"<a href='#{escape(t['id'])}'>{escape(t['id'])}</a>" for t in tables) + "</nav>",
]

for t in tables:
    tid = escape(t.get("id", ""))
    title = escape(t.get("title_zh", ""))
    notes = escape(t.get("notes", ""))
    cells = t.get("cells") or []
    parts.append(f"<h2 id='{tid}'>{tid} · {title}</h2>")
    parts.append(f"<p class='meta'>{notes}</p>")
    parts.append("<div class='pair'>")
    parts.append("<div><h3>汉字</h3><table>")
    for ri, row in enumerate(cells):
        parts.append("<tr>")
        tag = "th" if ri == 0 else "td"
        for cell in row:
            zh = escape(str(cell.get("zh", "")))
            parts.append(f"<{tag}>{zh}</{tag}>")
        parts.append("</tr>")
    parts.append("</table></div>")
    parts.append("<div><h3>盲文</h3><table class='braille'>")
    for ri, row in enumerate(cells):
        parts.append("<tr>")
        tag = "th" if ri == 0 else "td"
        for cell in row:
            br = escape(str(cell.get("braille", "")))
            parts.append(f"<{tag}>{br}</{tag}>")
        parts.append("</tr>")
    parts.append("</table></div></div>")

parts.append("</body></html>")
html_path = out_dir / "preview.html"
html_path.write_text("\n".join(parts), encoding="utf-8")
print("HTML", html_path, "bytes", html_path.stat().st_size)

wb = Workbook()
ws0 = wb.active
ws0.title = "目录"
ws0.append(["id", "标题", "行数", "列数", "说明"])
thin = Border(
    left=Side(style="thin", color="CCCCCC"),
    right=Side(style="thin", color="CCCCCC"),
    top=Side(style="thin", color="CCCCCC"),
    bottom=Side(style="thin", color="CCCCCC"),
)
header_fill = PatternFill("solid", fgColor="E8EEF8")
header_font = Font(bold=True)

for t in tables:
    tid = t.get("id", "")
    title = t.get("title_zh", "")
    cells = t.get("cells") or []
    nrows = len(cells)
    ncols = max((len(r) for r in cells), default=0)
    ws0.append([tid, title, nrows, ncols, t.get("notes", "")])

    name_zh = f"{tid}-汉字"[:31]
    wsz = wb.create_sheet(name_zh)
    wsz.append([f"{tid} {title} · 汉字"])
    for ri, row in enumerate(cells):
        vals = [c.get("zh", "") for c in row]
        wsz.append(vals)
        for ci in range(1, len(vals) + 1):
            cell = wsz.cell(ri + 2, ci)
            cell.border = thin
            cell.alignment = Alignment(wrap_text=True, vertical="center")
            if ri == 0:
                cell.fill = header_fill
                cell.font = header_font
    for col in wsz.columns:
        width = min(24, max(8, max(len(str(c.value or "")) for c in col) + 2))
        wsz.column_dimensions[col[0].column_letter].width = width

    name_br = f"{tid}-盲文"[:31]
    wsb = wb.create_sheet(name_br)
    wsb.append([f"{tid} {title} · 盲文"])
    for ri, row in enumerate(cells):
        vals = [c.get("braille", "") for c in row]
        wsb.append(vals)
        for ci in range(1, len(vals) + 1):
            cell = wsb.cell(ri + 2, ci)
            cell.border = thin
            cell.alignment = Alignment(wrap_text=True, vertical="center")
            if ri == 0:
                cell.fill = header_fill
                cell.font = header_font
    for col in wsb.columns:
        width = min(28, max(10, max(len(str(c.value or "")) for c in col) + 2))
        wsb.column_dimensions[col[0].column_letter].width = width

for col in ws0.columns:
    width = min(60, max(10, max(len(str(c.value or "")) for c in col) + 2))
    ws0.column_dimensions[col[0].column_letter].width = width
for cell in ws0[1]:
    cell.fill = header_fill
    cell.font = header_font

xlsx_path = out_dir / "preview.xlsx"
wb.save(xlsx_path)
print("XLSX", xlsx_path, "bytes", xlsx_path.stat().st_size)
print("TABLES", len(tables))
